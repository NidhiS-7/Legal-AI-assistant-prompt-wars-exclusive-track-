import io

import pytest
from docx import Document

from core.document_parser import (
    DocumentParseError,
    DocumentTooLargeError,
    UnsupportedFileTypeError,
    parse_document,
)


def test_parse_txt_document_success():
    content = b"This agreement is entered into by both parties."
    result = parse_document("lease.txt", content)
    assert "agreement" in result.text
    assert result.filename == "lease.txt"
    assert result.char_count == len(content)


def test_parse_rejects_oversized_file():
    content = b"a" * (2 * 1024 * 1024)
    with pytest.raises(DocumentTooLargeError):
        parse_document("big.txt", content, max_mb=1)


def test_parse_rejects_unsupported_extension():
    with pytest.raises(UnsupportedFileTypeError):
        parse_document("contract.exe", b"not a real binary")


def test_parse_rejects_file_with_no_extension():
    with pytest.raises(UnsupportedFileTypeError):
        parse_document("contract", b"some bytes")


def test_parse_rejects_empty_text_file():
    with pytest.raises(DocumentParseError):
        parse_document("empty.txt", b"   ")


def test_parse_docx_extracts_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Confidentiality Agreement")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "12 months"

    buffer = io.BytesIO()
    document.save(buffer)

    result = parse_document("agreement.docx", buffer.getvalue())
    assert "Confidentiality Agreement" in result.text
    assert "12 months" in result.text


def test_parse_handles_latin1_encoded_txt():
    content = "Café rental agreement".encode("latin-1")
    result = parse_document("lease.txt", content)
    assert "rental agreement" in result.text
